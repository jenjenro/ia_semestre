"""
Sistema de Análisis y Gestión Documental con Inteligencia Artificial
Semana 08 - Representaciones del Reconocimiento

Flujo completo:
Entrada (Documentos) -> Modelo de Reconocimiento (RNA) -> Evidencia (SQLite) -> Significado (Ontología GraphML)
"""

from datetime import datetime
from pathlib import Path
import os
import sqlite3
import unicodedata
import re

import docx
import joblib
import networkx as nx
import numpy as np
import pypdf
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix
from sklearn.model_selection import train_test_split
from sklearn.neural_network import MLPClassifier

# Rutas del proyecto
BASE_DIR = Path(__file__).resolve().parent.parent
SRC_DIR = BASE_DIR / "src"
DATA_DIR = BASE_DIR / "data"
ARTIFACTS_DIR = BASE_DIR / "artifacts"
REPORTS_DIR = BASE_DIR / "reports"
SAMPLES_DIR = ARTIFACTS_DIR / "muestras_prueba"

# Rutas de artefactos generados
MODEL_PATH = ARTIFACTS_DIR / "modelo_red_neuronal.joblib"
VECTORIZER_PATH = ARTIFACTS_DIR / "vectorizador_tfidf.joblib"
DB_PATH = ARTIFACTS_DIR / "evidencia_reconocimiento.db"
ONTOLOGY_PATH = ARTIFACTS_DIR / "ontologia_documental.graphml"
REPORT_PATH = REPORTS_DIR / "semana08.md"


# ==============================================================================
# 1. DATASET DEL DOMINIO DOCUMENTAL (60 Documentos etiquetados, 10 por clase)
# ==============================================================================

DATASET_DOCUMENTOS = [
    # 1. CONTRATOS (Dominio Legal)
    ("Contrato individual de trabajo a término indefinido celebrado entre la empresa y el empleado con cláusulas salariales y deberes.", "contrato"),
    ("Acuerdo de confidencialidad y no divulgación de secretos industriales y propiedad intelectual suscrito por las partes.", "contrato"),
    ("Contrato de prestación de servicios profesionales de consultoría tecnológica con cláusula penal, vigencia y firmas autorizadas.", "contrato"),
    ("Convenio interinstitucional de cooperación técnica y académica con compromisos bilaterales, pólizas de cumplimiento y representantes legales.", "contrato"),
    ("Contrato de arrendamiento de bien inmueble comercial con canon mensual, garantías, firmas autenticadas ante notaría y causales de terminación.", "contrato"),
    ("Contrato marco de suministro de mercancías y distribución comercial con cláusulas de exclusividad, resolución de controversias y firmas.", "contrato"),
    ("Acuerdo transaccional de terminación de controversias judiciales con renuncia expresa de acciones futuras y firmas de apoderados.", "contrato"),
    ("Contrato de confidencialidad para desarrollo conjunto y cesión de derechos patrimoniales de autor con firmas de las partes.", "contrato"),
    ("Contrato de compraventa mercantil de activos fijos con cláusula de saneamiento por evicción y entrega formal del inventario.", "contrato"),
    ("Acuerdo de nivel de servicio SLA contractual con penalizaciones por caída de plataforma y firmas de representantes autorizados.", "contrato"),

    # 2. FACTURAS (Dominio Financiero)
    ("Factura electrónica de venta número FE-8921 con detalle de ítems, impuesto al valor agregado IVA 19%, retención y total a pagar.", "factura"),
    ("Factura comercial por concepto de licencias de software en la nube, subtotal, descuento comercial, NIT emisor y código QR fiscal.", "factura"),
    ("Cuenta de cobro por servicios prestados con valor en letras, número de cuenta bancaria para consignación y firma del proveedor.", "factura"),
    ("Factura cambiaria de compraventa mercantil con fecha de vencimiento, forma de pago a crédito, tarifa de retefuente e impuestos aplicados.", "factura"),
    ("Recibo oficial de pago de obligaciones fiscales con discriminación de base gravable, tributos retenidos y firma de caja receptora.", "factura"),
    ("Factura electrónica de exportación en dólares con cotización TRM, partida arancelaria, datos del importador y liquidación de tributos.", "factura"),
    ("Factura de servicios públicos de energía y acueducto con consumo del periodo, lecturas de medidor, cargos fijos y fecha oportuna de pago.", "factura"),
    ("Factura de venta directa por honorarios profesionales con retención de ICA, IVA discriminado y pago mediante transferencia bancaria.", "factura"),
    ("Factura fiscal electrónica autorizada por la DIAN con resolución de facturación, total bruto gravable y valor neto a pagar.", "factura"),
    ("Comprobante de egreso y factura comercial por compra de suministros de oficina con tarifa impositiva y firma del cajero.", "factura"),

    # 3. REPORTES FINANCIEROS (Dominio Financiero)
    ("Reporte financiero anual con balance general consolidado, activos corrientes, pasivos exigibles y patrimonio neto de los accionistas.", "reporte_financiero"),
    ("Estado de resultados integral del tercer trimestre con ingresos operacionales, costo de ventas, margen EBITDA y utilidad neta.", "reporte_financiero"),
    ("Informe de auditoría contable y dictamen del revisor fiscal sobre los estados financieros y cumplimiento de normas contables NIIF.", "reporte_financiero"),
    ("Informe de ejecución presupuestal con análisis de desviaciones de gasto, proyecciones de flujo de caja y rentabilidad de inversiones.", "reporte_financiero"),
    ("Estado de cambios en la situación financiera con origen y aplicación de recursos, amortizaciones y pasivos a largo plazo.", "reporte_financiero"),
    ("Resumen ejecutivo de rendición de cuentas financieras con indicadores de liquidez, solvencia, endeudamiento y rotación de cartera.", "reporte_financiero"),
    ("Informe pericial contable de valoración económica de la empresa con flujo de fondos descontados y múltiplos financieros comparables.", "reporte_financiero"),
    ("Informe de flujo de efectivo trimestral con actividades de operación, inversión de capital y financiación empresarial.", "reporte_financiero"),
    ("Dictamen de estados financieros por contador público titulado con balance general y notas explicativas a los estados contables.", "reporte_financiero"),
    ("Análisis financiero de razones de endeudamiento, capital de trabajo neto operativo y margen de utilidad antes de impuestos.", "reporte_financiero"),

    # 4. MANUALES TÉCNICOS (Dominio Tecnológico)
    ("Manual técnico de instalación y despliegue de arquitectura de microservicios con Docker, Kubernetes y configuración de variables de entorno.", "manual_tecnico"),
    ("Guía técnica de desarrollo e integración de API RESTful con autenticación OAuth2, especificación OpenAPI endpoints y control de excepciones.", "manual_tecnico"),
    ("Manual de usuario y administración del sistema de análisis documental con especificaciones de base de datos y puertos de red.", "manual_tecnico"),
    ("Documentación de arquitectura técnica de software con diagramas de componentes, patrones de diseño y requerimientos de hardware.", "manual_tecnico"),
    ("Manual de procedimientos de mantenimiento preventivo, copias de seguridad de bases de datos PostgreSQL y políticas de recuperación ante desastres.", "manual_tecnico"),
    ("Guía de configuración de cortafuegos de red, cifrado TLS 1.3, gestión de certificados SSL y monitoreo de tráfico con Prometheus.", "manual_tecnico"),
    ("Manual de ingeniería de datos para pipelines ETL en Apache Spark con transformación de datasets y almacenamiento en data lakehouse.", "manual_tecnico"),
    ("Guía de referencia de arquitectura en la nube con políticas de balanceo de carga, VPC y despliegue continuo en pipelines CI/CD.", "manual_tecnico"),
    ("Manual técnico de configuración de servidor web Nginx, configuración de proxy inverso y protección con certificados SSL TLS.", "manual_tecnico"),
    ("Especificación técnica de requisitos de software con casos de uso, diagramas entidad relación de base de datos relacional y APIs.", "manual_tecnico"),

    # 5. MEMORANDOS (Dominio Administrativo)
    ("Memorando interno dirigido a todo el personal sobre actualización de políticas de trabajo remoto, teletrabajo y registro de jornada laboral.", "memorando"),
    ("Comunicado interno de gerencia sobre las fechas límites de reporte de vacaciones anuales y procedimiento ante recursos humanos.", "memorando"),
    ("Circular interna informativa comunicando cambios en la estructura organizacional, nuevos nombramientos y líneas de reporte directo.", "memorando"),
    ("Memorando de amonestación y llamado de atención disciplinario por reiteradas ausencias injustificadas y violación del reglamento interno.", "memorando"),
    ("Aviso administrativo institucional sobre el cronograma de evaluación de desempeño laboral y capacitaciones obligatorias de la empresa.", "memorando"),
    ("Circular general sobre protocolos de seguridad y salud en el trabajo, uso de dotación y reporte de condiciones de riesgo laboral.", "memorando"),
    ("Memorando de asignación de responsabilidades para el comité de compras extraordinarias y adquisición de equipos de cómputo.", "memorando"),
    ("Memorando de talento humano sobre el cronograma oficial de días festivos corporativos y turnos especiales de trabajo programados.", "memorando"),
    ("Circular informativa de gerencia general recordando los lineamientos éticos del código de conducta corporativa y transparencia.", "memorando"),
    ("Comunicado institucional sobre el protocolo formal de atención y solicitud anticipada de permisos laborales remunerados.", "memorando"),

    # 6. ACTAS DE REUNIÓN (Dominio Administrativo)
    ("Acta de reunión ordinaria de la junta directiva número 45 con verificación de quórum, lectura del orden del día y votación de acuerdos.", "acta_reunion"),
    ("Acta de asamblea general ordinaria de accionistas con deliberaciones sobre aprobación de dividendos, balance general y firmas de comisión escrutadora.", "acta_reunion"),
    ("Minuta de reunión de comité técnico de desarrollo de software con asignación de compromisos, fechas de entrega y participantes convocados.", "acta_reunion"),
    ("Acta de reunión del comité de convivencia laboral con análisis de casos presentados, compromisos de mediación y firmas de los asistentes.", "acta_reunion"),
    ("Minuta de seguimiento de proyecto de implementación tecnológica con estado de tareas, riesgos identificados y acuerdos entre líderes de área.", "acta_reunion"),
    ("Acta de posesión de cargo administrativo con juramento de deberes legales, presentación de documentos de identidad y constancia escrita.", "acta_reunion"),
    ("Acta formal de cierre de proyecto con evaluación de entregables, cumplimiento de cronograma, lecciones aprendidas y visto bueno final.", "acta_reunion"),
    ("Acta de comité de dirección semanal con revisión detallada de indicadores clave de desempeño KPI y compromisos estratégicos adquiridos.", "acta_reunion"),
    ("Minuta de sesión de comité de seguridad de la información con acuerdos formales sobre políticas de control de accesos y auditoría.", "acta_reunion"),
    ("Acta solemne de junta directiva donde se aprueba por decisión unánime el nombramiento del nuevo auditor interno institucional.", "acta_reunion"),
]


# ==============================================================================
# 2. EXTRACCIÓN Y LECTURA MULTIFORMATO DE DOCUMENTOS
# ==============================================================================

def extraer_texto_documento(fuente) -> tuple[str, str, str]:
    """
    Extrae texto de documentos en múltiples formatos (.txt, .md, .pdf, .docx)
    o procesa directamente cadenas de texto plano.
    
    Retorna:
        tuple[nombre_fuente, texto_extraido, tipo_fuente]
    """
    try:
        ruta = Path(fuente)
    except Exception:
        ruta = None

    if ruta is None or not ruta.exists() or not ruta.is_file():
        texto_str = str(fuente).strip()
        nombre = "texto_directo" if len(texto_str) > 30 else f"muestra_{texto_str[:15].replace(' ', '_')}"
        return nombre, texto_str, "texto_directo"

    nombre = ruta.name
    sufijo = ruta.suffix.lower()
    texto = ""

    if sufijo in [".txt", ".md"]:
        try:
            texto = ruta.read_text(encoding="utf-8", errors="ignore")
        except Exception as e:
            texto = f"[Error lectura texto]: {e}"
    elif sufijo == ".pdf":
        try:
            lector = pypdf.PdfReader(str(ruta))
            texto = "\n".join([pag.extract_text() or "" for pag in lector.pages])
        except Exception as e:
            print(f"[!] Error leyendo PDF {nombre}: {e}")
            texto = ""
    elif sufijo == ".docx":
        try:
            doc = docx.Document(str(ruta))
            texto = "\n".join([p.text for p in doc.paragraphs if p.text])
        except Exception as e:
            print(f"[!] Error leyendo Word {nombre}: {e}")
            texto = ""
    else:
        try:
            texto = ruta.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            texto = ""

    return nombre, texto.strip(), sufijo.replace(".", "") or "binario"


def crear_muestras_prueba() -> list[tuple[Path, str]]:
    """Crea archivos reales en disco (.txt, .md, .docx) para demostrar la ingesta multiformato."""
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    muestras = []

    # 1. Contrato (.txt)
    p_contrato = SAMPLES_DIR / "contrato_servicios_2026.txt"
    p_contrato.write_text(
        "CONTRATO DE PRESTACIÓN DE SERVICIOS TÉCNICOS\n"
        "Entre Soluciones Digitales S.A.S. y el consultor independiente. "
        "Cláusula primera: Objeto del contrato. Cláusula segunda: Honorarios pactados. "
        "Cláusula penal por incumplimiento y deber estricto de confidencialidad. "
        "Firmas autorizadas de las partes suscritas ante notario.",
        encoding="utf-8"
    )
    muestras.append((p_contrato, "contrato"))

    # 2. Factura (.txt)
    p_factura = SAMPLES_DIR / "factura_proveedor_servicios.txt"
    p_factura.write_text(
        "FACTURA ELECTRÓNICA DE VENTA FE-44910\n"
        "Emisor: Infraestructura Cloud Andina S.A.S. - NIT: 900.542.118-4\n"
        "Cliente: Corporación Universitaria\n"
        "Concepto: Servicios de computación en la nube periodo agosto 2026\n"
        "Subtotal: $ 4.500.000 COP | IVA 19%: $ 855.000 COP | Retención Fuente: $ 180.000 COP\n"
        "Total neto a pagar: $ 5.175.000 COP. Código QR fiscal DIAN.",
        encoding="utf-8"
    )
    muestras.append((p_factura, "factura"))

    # 3. Reporte Financiero (.docx)
    p_reporte = SAMPLES_DIR / "balance_general_ejercicio.docx"
    doc_word = docx.Document()
    doc_word.add_heading("INFORME DE AUDITORÍA Y BALANCE GENERAL FINANCIERO", level=1)
    doc_word.add_paragraph(
        "El presente estado financiero consolidado contiene los activos corrientes, pasivos exigibles, "
        "y el patrimonio neto de la organización. Según el dictamen pericial del revisor fiscal, "
        "los estados de resultados integrales y flujos de efectivo cumplen con las normas NIIF."
    )
    doc_word.save(str(p_reporte))
    muestras.append((p_reporte, "reporte_financiero"))

    # 4. Manual Técnico (.md)
    p_manual = SAMPLES_DIR / "manual_despliegue_contenedores.md"
    p_manual.write_text(
        "# Manual Técnico de Despliegue de Microservicios\n\n"
        "Este documento describe la arquitectura técnica de contenedores Docker y orquestación con Kubernetes. "
        "Configure las variables de entorno, exponga los endpoints de la API RESTful mediante ingress controller "
        "y verifique los registros de logs y puertos de red correspondientes.",
        encoding="utf-8"
    )
    muestras.append((p_manual, "manual_tecnico"))

    # 5. Memorando (.txt)
    p_memo = SAMPLES_DIR / "circular_politica_laboral.txt"
    p_memo.write_text(
        "MEMORANDO INTERNO DE TALENTO HUMANO\n"
        "Para: Todos los colaboradores de planta\n"
        "De: Dirección de Gestión Humana\n"
        "Asunto: Actualización de políticas de trabajo remoto y registro de jornada laboral.\n"
        "Se informa a todo el personal que las solicitudes de permisos remunerados deben radicarse "
        "con al menos tres días de anticipación según el reglamento interno.",
        encoding="utf-8"
    )
    muestras.append((p_memo, "memorando"))

    # 6. Acta de Reunión (.txt)
    p_acta = SAMPLES_DIR / "acta_comite_directivo.txt"
    p_acta.write_text(
        "ACTA DE REUNIÓN ORDINARIA DE JUNTA DIRECTIVA N° 12\n"
        "En la ciudad de Bogotá, se reúnen los miembros convocados previa verificación del quórum legal. "
        "Lectura del orden del día. Se aprueban por unanimidad los puntos tratados y se fijan compromisos "
        "con plazos específicos. Firman el presidente y el secretario de la sesión.",
        encoding="utf-8"
    )
    muestras.append((p_acta, "acta_reunion"))

    # 7. Documento Atípico / Fuera de Dominio (.txt)
    p_receta = SAMPLES_DIR / "receta_reposteria_atipica.txt"
    p_receta.write_text(
        "RECETA DE COCINA CASERA Y REPOSTERÍA\n"
        "Ingredientes: 500 gramos de harina de trigo, 3 huevos frescos, 200 gramos de mantequilla y azúcar. "
        "Preparación: Mezclar en un tazón hasta obtener una masa homogénea y hornear a 180 grados centígrados "
        "durante cuarenta y cinco minutos.",
        encoding="utf-8"
    )
    muestras.append((p_receta, "desconocido"))

    return muestras


# ==============================================================================
# 3. BASE DE DATOS SQLITE (REGISTRO DE EVIDENCIA)
# ==============================================================================

def inicializar_base_datos(db_path: Path = DB_PATH) -> None:
    """Crea la tabla de evidencia relacional en SQLite con sus respectivos índices."""
    db_path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS evidencia_reconocimiento (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT NOT NULL,
                fuente_documento TEXT NOT NULL,
                tipo_fuente TEXT NOT NULL,
                resumen_contenido TEXT NOT NULL,
                categoria_real TEXT NOT NULL,
                categoria_predicha TEXT NOT NULL,
                nivel_confianza REAL NOT NULL,
                estado_proceso TEXT NOT NULL,
                categoria_dominio TEXT NOT NULL,
                departamento_destino TEXT NOT NULL,
                accion_sugerida TEXT NOT NULL,
                significado_ontologico TEXT NOT NULL
            );
        """)
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_cat_predicha ON evidencia_reconocimiento(categoria_predicha);")
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_estado ON evidencia_reconocimiento(estado_proceso);")
        conn.commit()


def registrar_evidencia(registro: dict, db_path: Path = DB_PATH) -> int:
    """Inserta un registro de evidencia en SQLite y retorna su ID autoincremental."""
    with sqlite3.connect(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO evidencia_reconocimiento (
                timestamp,
                fuente_documento,
                tipo_fuente,
                resumen_contenido,
                categoria_real,
                categoria_predicha,
                nivel_confianza,
                estado_proceso,
                categoria_dominio,
                departamento_destino,
                accion_sugerida,
                significado_ontologico
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, (
            registro.get("timestamp", datetime.now().isoformat()),
            registro.get("fuente_documento", "desconocido"),
            registro.get("tipo_fuente", "texto"),
            registro.get("resumen_contenido", "")[:250],
            registro.get("categoria_real", "No especificada"),
            registro.get("categoria_predicha", "desconocido"),
            float(registro.get("nivel_confianza", 0.0)),
            registro.get("estado_proceso", "Procesado"),
            registro.get("categoria_dominio", "General"),
            registro.get("departamento_destino", "Revisión General"),
            registro.get("accion_sugerida", "Archivo"),
            registro.get("significado_ontologico", "Sin interpretación")
        ))
        conn.commit()
        return cursor.lastrowid


def consultar_evidencias(db_path: Path = DB_PATH, limite: int = 15) -> list[dict]:
    """Consulta los registros de evidencia almacenados en la base de datos SQLite."""
    with sqlite3.connect(db_path) as conn:
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        cursor.execute("""
            SELECT id, timestamp, fuente_documento, tipo_fuente, categoria_real, categoria_predicha,
                   nivel_confianza, estado_proceso, categoria_dominio, departamento_destino
            FROM evidencia_reconocimiento
            ORDER BY id DESC
            LIMIT ?;
        """, (limite,))
        rows = cursor.fetchall()
        return [dict(r) for r in rows]


def consultar_estadisticas_bd(db_path: Path = DB_PATH) -> dict:
    """Obtiene métricas agregadas de los registros de evidencia almacenados en SQLite."""
    with sqlite3.connect(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM evidencia_reconocimiento;")
        total = cursor.fetchone()[0]

        cursor.execute("SELECT AVG(nivel_confianza) FROM evidencia_reconocimiento WHERE nivel_confianza > 0;")
        prom_conf = cursor.fetchone()[0] or 0.0

        cursor.execute("SELECT categoria_predicha, COUNT(*) FROM evidencia_reconocimiento GROUP BY categoria_predicha;")
        por_cat = dict(cursor.fetchall())

        cursor.execute("SELECT estado_proceso, COUNT(*) FROM evidencia_reconocimiento GROUP BY estado_proceso;")
        por_estado = dict(cursor.fetchall())

        return {
            "total_registros": total,
            "confianza_promedio": prom_conf,
            "distribucion_categorias": por_cat,
            "distribucion_estados": por_estado
        }


# ==============================================================================
# 4. CONSTRUCCIÓN DE LA ONTOLOGÍA (GRAPHML)
# ==============================================================================

def construir_ontologia() -> nx.DiGraph:
    """
    Construye la ontología del dominio del proyecto en forma de grafo dirigido (NetworkX).
    Incluye más de 5 conceptos principales, subconceptos del dominio y relaciones con sentido de frase:
    Sujeto -> Predicado (etiqueta) -> Objeto.
    """
    grafo = nx.DiGraph()

    # Conceptos Fundamentales del Sistema Documental (Nivel Superior)
    conceptos_nucleo = {
        "DocumentoDigital": {"tipo": "Entidad", "descripcion": "Archivo digital o texto ingresado al sistema documental."},
        "RedNeuronal": {"tipo": "AgenteIA", "descripcion": "Modelo de Perceptrón Multicapa (MLP) que procesa patrones y genera inferencias."},
        "Prediccion": {"tipo": "Inferencia", "descripcion": "Resultado categórico probabilístico asignado a un documento."},
        "EvidenciaBD": {"tipo": "Almacenamiento", "descripcion": "Registro persistente en base de datos SQLite con trazabilidad del procesamiento."},
        "TipoDocumento": {"tipo": "Clasificacion", "descripcion": "Taxonomía de clases documentales del sistema."},
        "CategoriaDominio": {"tipo": "Dominio", "descripcion": "Área macro de negocio u organizacional a la que pertenece el documento."},
        "DepartamentoResponsable": {"tipo": "UnidadOrganizacional", "descripcion": "Dependencia encargada de la gestión y custodia del documento."},
        "AccionFlujo": {"tipo": "PoliticaOperativa", "descripcion": "Procedimiento o acción inmediata requerida según el tipo de documento."}
    }

    for nodo, attrs in conceptos_nucleo.items():
        grafo.add_node(nodo, **attrs)

    # Subconceptos específicos del dominio (Instancias taxonómicas)
    subconceptos = {
        # Tipos documentales
        "contrato": {"tipo": "SubclaseTipo", "nombre_legible": "Contrato Legal"},
        "factura": {"tipo": "SubclaseTipo", "nombre_legible": "Factura Comercial"},
        "reporte_financiero": {"tipo": "SubclaseTipo", "nombre_legible": "Reporte Financiero"},
        "manual_tecnico": {"tipo": "SubclaseTipo", "nombre_legible": "Manual Técnico"},
        "memorando": {"tipo": "SubclaseTipo", "nombre_legible": "Memorando Interno"},
        "acta_reunion": {"tipo": "SubclaseTipo", "nombre_legible": "Acta de Reunión"},

        # Dominios organizacionales
        "Dominio_Legal": {"tipo": "SubclaseDominio", "area": "Asuntos Jurídicos y Normativos"},
        "Dominio_Financiero": {"tipo": "SubclaseDominio", "area": "Gestión Contable, Tributaria y Fiscal"},
        "Dominio_Tecnologico": {"tipo": "SubclaseDominio", "area": "Infraestructura, TI y Desarrollo"},
        "Dominio_Administrativo": {"tipo": "SubclaseDominio", "area": "Gestión Humana y Gobierno Corporativo"},

        # Departamentos responsables
        "Area_Juridica": {"tipo": "Departamento", "oficina": "Gerencia Jurídica y Cumplimiento"},
        "Contabilidad_Tesoreria": {"tipo": "Departamento", "oficina": "Departamento de Finanzas y Contabilidad"},
        "Operaciones_TI": {"tipo": "Departamento", "oficina": "Dirección de Tecnología de la Información"},
        "Gestion_Humana": {"tipo": "Departamento", "oficina": "Dirección de Talento Humano y Operaciones"},

        # Acciones operativas
        "Custodia_Firmas_Legales": {"tipo": "Accion", "descripcion": "Verificación de personería jurídica, autenticidad de firmas y custodia bajo bóveda."},
        "Auditoria_Tributaria": {"tipo": "Accion", "descripcion": "Conciliación de impuestos, liquidación de retenciones y registro en libro mayor."},
        "Despliegue_Documentacion_Tecnica": {"tipo": "Accion", "descripcion": "Publicación en repositorio técnico de ingeniería y validación de estándares de código."},
        "Difusion_Interna": {"tipo": "Accion", "descripcion": "Notificación a colaboradores y registro en actas de seguimiento institucional."}
    }

    for nodo, attrs in subconceptos.items():
        grafo.add_node(nodo, **attrs)

    # Relaciones entre conceptos núcleo (frases semánticas principales: Sujeto -> Predicado -> Objeto)
    relaciones_nucleo = [
        ("RedNeuronal", "DocumentoDigital", "analiza_patrones_de"),
        ("RedNeuronal", "Prediccion", "genera_prediccion"),
        ("Prediccion", "TipoDocumento", "asigna_clase_a"),
        ("Prediccion", "EvidenciaBD", "registra_evidencia_en"),
        ("EvidenciaBD", "DocumentoDigital", "garantiza_trazabilidad_de"),
        ("TipoDocumento", "CategoriaDominio", "pertenece_a_dominio"),
        ("TipoDocumento", "DepartamentoResponsable", "enruta_hacia"),
        ("TipoDocumento", "AccionFlujo", "determina_accion")
    ]

    for origen, destino, relacion in relaciones_nucleo:
        grafo.add_edge(origen, destino, label=relacion, relacion_frase=f"{origen} {relacion} {destino}")

    # Relaciones de clasificación e inferencia específicas
    mapeos_especificos = [
        # Contrato
        ("TipoDocumento", "contrato", "incluye_tipo"),
        ("contrato", "Dominio_Legal", "pertenece_a_dominio"),
        ("contrato", "Area_Juridica", "enruta_hacia"),
        ("contrato", "Custodia_Firmas_Legales", "determina_accion"),

        # Factura
        ("TipoDocumento", "factura", "incluye_tipo"),
        ("factura", "Dominio_Financiero", "pertenece_a_dominio"),
        ("factura", "Contabilidad_Tesoreria", "enruta_hacia"),
        ("factura", "Auditoria_Tributaria", "determina_accion"),

        # Reporte Financiero
        ("TipoDocumento", "reporte_financiero", "incluye_tipo"),
        ("reporte_financiero", "Dominio_Financiero", "pertenece_a_dominio"),
        ("reporte_financiero", "Contabilidad_Tesoreria", "enruta_hacia"),
        ("reporte_financiero", "Auditoria_Tributaria", "determina_accion"),

        # Manual Técnico
        ("TipoDocumento", "manual_tecnico", "incluye_tipo"),
        ("manual_tecnico", "Dominio_Tecnologico", "pertenece_a_dominio"),
        ("manual_tecnico", "Operaciones_TI", "enruta_hacia"),
        ("manual_tecnico", "Despliegue_Documentacion_Tecnica", "determina_accion"),

        # Memorando
        ("TipoDocumento", "memorando", "incluye_tipo"),
        ("memorando", "Dominio_Administrativo", "pertenece_a_dominio"),
        ("memorando", "Gestion_Humana", "enruta_hacia"),
        ("memorando", "Difusion_Interna", "determina_accion"),

        # Acta de Reunión
        ("TipoDocumento", "acta_reunion", "incluye_tipo"),
        ("acta_reunion", "Dominio_Administrativo", "pertenece_a_dominio"),
        ("acta_reunion", "Gestion_Humana", "enruta_hacia"),
        ("acta_reunion", "Difusion_Interna", "determina_accion"),
    ]

    for origen, destino, relacion in mapeos_especificos:
        grafo.add_edge(origen, destino, label=relacion, relacion_frase=f"{origen} {relacion} {destino}")

    return grafo


def guardar_ontologia_graphml(grafo: nx.DiGraph, ruta_archivo: Path = ONTOLOGY_PATH) -> None:
    """Exporta el grafo ontológico al formato estándar GraphML."""
    ruta_archivo.parent.mkdir(parents=True, exist_ok=True)
    nx.write_graphml(grafo, str(ruta_archivo))


def interpretar_con_ontologia(ontologia: nx.DiGraph, clase_predicha: str) -> dict:
    """
    Navega el grafo ontológico a partir de la clase predicha para extraer su significado:
    dominio macro, departamento asignado y acción operativa.
    """
    clase_predicha_norm = clase_predicha.lower().strip()
    
    if clase_predicha_norm not in ontologia:
        return {
            "categoria_dominio": "Dominio_Desconocido",
            "departamento_destino": "Mesa_Entrada_General",
            "accion_sugerida": "Revision_Manual_Falta_De_Regla",
            "significado_ontologico": f"El concepto '{clase_predicha}' no posee relación formal dentro de la ontología del proyecto."
        }

    # Búsqueda de relaciones salientes
    dominio = "General"
    departamento = "Administracion"
    accion = "Archivo_General"

    for _, destino, datos in ontologia.out_edges(clase_predicha_norm, data=True):
        rel = datos.get("label", "")
        if rel == "pertenece_a_dominio":
            dominio = destino
        elif rel == "enruta_hacia":
            departamento = destino
        elif rel == "determina_accion":
            accion = destino

    significado = (
        f"El documento fue categorizado como '{clase_predicha_norm}', perteneciente a '{dominio}'. "
        f"Se enruta a '{departamento}' con instrucción de aplicar '{accion}'."
    )

    return {
        "categoria_dominio": dominio,
        "departamento_destino": departamento,
        "accion_sugerida": accion,
        "significado_ontologico": significado
    }


# ==============================================================================
# 5. ENTRENAMIENTO Y VALIDACIÓN DEL MODELO NEURONAL
# ==============================================================================

def entrenar_modelo_neuronal(random_state: int = 42) -> tuple[MLPClassifier, TfidfVectorizer, dict]:
    """
    Entrena una Red Neuronal Artificial (Multi-Layer Perceptron) con TF-IDF.
    Realiza partición estratificada de entrenamiento y prueba, calcula métricas de validación
    y guarda los artefactos en disco.
    """
    textos = [elem[0] for elem in DATASET_DOCUMENTOS]
    etiquetas = [elem[1] for elem in DATASET_DOCUMENTOS]

    # Partición estratificada (80% entrenamiento: 48 docs, 20% prueba: 12 docs)
    X_train_raw, X_test_raw, y_train, y_test = train_test_split(
        textos,
        etiquetas,
        test_size=0.20,
        random_state=random_state,
        stratify=etiquetas
    )

    # Vectorización TF-IDF con unigramas y bigramas
    vectorizador = TfidfVectorizer(
        ngram_range=(1, 2),
        sublinear_tf=True,
        max_features=1000
    )
    X_train = vectorizador.fit_transform(X_train_raw)
    X_test = vectorizador.transform(X_test_raw)

    # Red Neuronal Artificial (Perceptrón Multicapa)
    # Capa oculta con 32 neuronas artificiales con activación ReLU y optimizador Adam
    modelo = MLPClassifier(
        hidden_layer_sizes=(32,),
        activation="relu",
        solver="adam",
        alpha=0.001,
        learning_rate_init=0.005,
        max_iter=600,
        random_state=random_state,
        early_stopping=False
    )
    modelo.fit(X_train, y_train)

    # Validación con el conjunto de prueba
    y_pred = modelo.predict(X_test)
    y_proba = modelo.predict_proba(X_test)
    accuracy = accuracy_score(y_test, y_pred)
    reporte_dict = classification_report(y_test, y_pred, output_dict=True, zero_division=0)
    matriz_conf = confusion_matrix(y_test, y_pred, labels=modelo.classes_)

    # Guardado de artefactos
    ARTIFACTS_DIR.mkdir(parents=True, exist_ok=True)
    joblib.dump(modelo, MODEL_PATH)
    joblib.dump(vectorizador, VECTORIZER_PATH)

    resultados_evaluacion = {
        "accuracy": accuracy,
        "classification_report": reporte_dict,
        "confusion_matrix": matriz_conf,
        "classes": list(modelo.classes_),
        "n_train": len(X_train_raw),
        "n_test": len(X_test_raw),
        "y_test": y_test,
        "y_pred": list(y_pred),
        "loss_curve_final": modelo.loss_,
        "n_iter": modelo.n_iter_
    }

    return modelo, vectorizador, resultados_evaluacion


# ==============================================================================
# 6. PIPELINE DE INTEGRACIÓN: ENTRADA -> RNA -> BD -> ONTOLOGÍA
# ==============================================================================

def clasificar_y_registrar_documento(
    fuente,
    modelo: MLPClassifier,
    vectorizador: TfidfVectorizer,
    ontologia: nx.DiGraph,
    db_path: Path = DB_PATH,
    categoria_real: str = "No especificada",
    nombre_documento: str = None,
    umbral_minimo: float = 0.40
) -> dict:
    """
    Ejecuta el flujo completo:
    1. Entrada: Extrae texto de archivo (.txt, .md, .pdf, .docx) o texto plano.
    2. RNA: Extrae características TF-IDF y predice categoría y confianza.
    3. Ontología: Interpreta el significado del documento (dominio, oficina, acción).
    4. BD: Registra la evidencia en SQLite con sello de tiempo.
    """
    nombre_extraido, texto, tipo_fuente = extraer_texto_documento(fuente)
    nombre = nombre_documento or nombre_extraido

    # Validación 1: Documento con texto legible
    if not texto:
        resultado = {
            "timestamp": datetime.now().isoformat(),
            "fuente_documento": nombre,
            "tipo_fuente": tipo_fuente,
            "resumen_contenido": "[Vacío o no legible]",
            "categoria_real": categoria_real,
            "categoria_predicha": "Ilegible",
            "nivel_confianza": 0.0,
            "estado_proceso": "Rechazado (Sin contenido textual legible o requiere OCR)",
            "categoria_dominio": "No_Aplica",
            "departamento_destino": "Mesa_Ayuda_Digitalizacion",
            "accion_sugerida": "Escanear_Nuevamente_Con_OCR",
            "significado_ontologico": "El documento no pudo ser procesado porque carece de capa de texto."
        }
        id_db = registrar_evidencia(resultado, db_path)
        resultado["id_bd"] = id_db
        return resultado

    # Transformación a vector de patrones
    vector = vectorizador.transform([texto])

    # Validación 2: Coincidencia de vocabulario del dominio
    if vector.sum() == 0:
        resultado = {
            "timestamp": datetime.now().isoformat(),
            "fuente_documento": nombre,
            "tipo_fuente": tipo_fuente,
            "resumen_contenido": texto[:200],
            "categoria_real": categoria_real,
            "categoria_predicha": "Fuera_De_Dominio",
            "nivel_confianza": 0.0,
            "estado_proceso": "Rechazado (Vocabulario desconocido)",
            "categoria_dominio": "No_Identificada",
            "departamento_destino": "Revision_Manual_Externa",
            "accion_sugerida": "Analisis_Por_Especialista",
            "significado_ontologico": "El vocabulario del documento no contiene patrones coincidentes con ninguna categoría conocida del repositorio."
        }
        id_db = registrar_evidencia(resultado, db_path)
        resultado["id_bd"] = id_db
        return resultado

    # Inferencia con la Red Neuronal
    probabilidades = modelo.predict_proba(vector)[0]
    idx_max = int(probabilidades.argmax())
    confianza = float(probabilidades[idx_max])
    clase_predicha = str(modelo.classes_[idx_max])

    # Interpretación Ontológica
    interpretacion = interpretar_con_ontologia(ontologia, clase_predicha)

    # Evaluación de certidumbre
    if confianza < umbral_minimo:
        estado = f"Requiere Revisión Manual (Confianza {confianza*100:.1f}% < {umbral_minimo*100:.0f}%)"
    else:
        estado = "Procesado Exitosamente"

    resultado = {
        "timestamp": datetime.now().isoformat(),
        "fuente_documento": nombre,
        "tipo_fuente": tipo_fuente,
        "resumen_contenido": texto[:200].replace("\n", " "),
        "categoria_real": categoria_real,
        "categoria_predicha": clase_predicha,
        "nivel_confianza": round(confianza, 4),
        "estado_proceso": estado,
        "categoria_dominio": interpretacion["categoria_dominio"],
        "departamento_destino": interpretacion["departamento_destino"],
        "accion_sugerida": interpretacion["accion_sugerida"],
        "significado_ontologico": interpretacion["significado_ontologico"]
    }

    # Registro en Base de Datos SQLite (Evidencia)
    id_db = registrar_evidencia(resultado, db_path)
    resultado["id_bd"] = id_db

    return resultado


# ==============================================================================
# 7. GENERACIÓN AUTOMÁTICA DEL REPORTE semana08.md
# ==============================================================================

def generar_reporte_md(evaluacion: dict, muestras_procesadas: list[dict], ontologia: nx.DiGraph) -> None:
    """Genera el reporte Markdown semana08.md cumpliendo todos los requerimientos."""
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)

    acc_pct = evaluacion["accuracy"] * 100
    classes = evaluacion["classes"]
    rep = evaluacion["classification_report"]

    tabla_metricas = "| Categoría | Precisión | Recall | F1-Score | Muestras Test |\n| :--- | :---: | :---: | :---: | :---: |\n"
    for c in classes:
        if c in rep:
            p = rep[c]["precision"] * 100
            r = rep[c]["recall"] * 100
            f = rep[c]["f1-score"] * 100
            s = int(rep[c]["support"])
            tabla_metricas += f"| **{c}** | {p:.1f}% | {r:.1f}% | {f:.1f}% | {s} |\n"
    tabla_metricas += f"| **Promedio Ponderado** | {rep['weighted avg']['precision']*100:.1f}% | {rep['weighted avg']['recall']*100:.1f}% | {rep['weighted avg']['f1-score']*100:.1f}% | {evaluacion['n_test']} |\n"

    # Matriz de confusión
    header_matriz = "| Real \\ Pred | " + " | ".join([f"**{c[:6]}**" for c in classes]) + " |\n"
    sep_matriz = "| :--- | " + " | ".join([":---:" for _ in classes]) + " |\n"
    filas_matriz = ""
    for i, c in enumerate(classes):
        fila_valores = " | ".join([str(val) for val in evaluacion["confusion_matrix"][i]])
        filas_matriz += f"| **{c[:6]}** | {fila_valores} |\n"
    tabla_confusion = header_matriz + sep_matriz + filas_matriz

    # Relaciones de la ontología
    relaciones_md = ""
    for u, v, d in ontologia.edges(data=True):
        rel = d.get("label", "relacionado_con")
        relaciones_md += f"- `{u}` **{rel}** `{v}`\n"

    # Casos de prueba procesados
    tabla_casos = "| ID BD | Documento Analizado | Formato | Cat. Real | Predicción RNA | Confianza | Estado | Destino Ontológico |\n| :---: | :--- | :---: | :--- | :--- | :---: | :--- | :--- |\n"
    for m in muestras_procesadas:
        nombre = m["fuente_documento"][:35]
        fmt = m.get("tipo_fuente", "txt")
        c_real = m["categoria_real"]
        c_pred = m["categoria_predicha"]
        conf = f"{m['nivel_confianza']*100:.1f}%"
        estado = "Procesado" if "Procesado" in m["estado_proceso"] else "Revisión"
        destino = f"{m['departamento_destino']} ({m['categoria_dominio']})"
        tabla_casos += f"| {m['id_bd']} | `{nombre}` | `{fmt}` | {c_real} | **{c_pred}** | {conf} | {estado} | {destino} |\n"

    contenido = f"""# Reporte Semana 08 - Representaciones del Reconocimiento

**Proyecto:** Sistema de Análisis y Gestión Documental con Inteligencia Artificial  
**Fecha de corte:** Septiembre 2026  
**Modalidad:** Individual  
**Entrega:** Repositorio GitHub – Commit Semana 8  

---

## 1. Qué reconoce el sistema

El sistema inteligente implementado aborda el reconocimiento y clasificación automática de documentos organizacionales y administrativos dentro del flujo documental institucional. En concordancia con el proyecto desarrollado durante el semestre, el componente inteligente reconoce **seis categorías documentales fundamentales**:

1. **`contrato`**: Acuerdos de voluntades, contratos laborales, de prestación de servicios, convenios y acuerdos de confidencialidad (NDA).
2. **`factura`**: Facturas electrónicas de venta, cuentas de cobro, comprobantes fiscales y recibos de pago.
3. **`reporte_financiero`**: Balances generales, estados de resultados, informes de auditoría contable y ejecuciones presupuestales.
4. **`manual_tecnico`**: Guías de instalación, especificaciones de arquitectura de software, manuales de usuario y documentación de APIs.
5. **`memorando`**: Comunicaciones internas de talento humano, circulares informativas y avisos institucionales.
6. **`acta_reunion`**: Actas de junta directiva, asambleas de accionistas, comités técnicos y minutas de seguimiento de proyectos.

El sistema recibe documentos en texto plano o archivos digitales en formatos **`.txt`, `.md`, `.pdf` y `.docx`**, reconociendo su contenido textual y determinando la tipología documental correspondiente.

---

## 2. Cómo funciona el modelo de reconocimiento (RNA)

### Arquitectura de la Red Neuronal Artificial
El reconocimiento de patrones no se realiza mediante palabras clave fijas o reglas rígidas, sino a través de un **Perceptrón Multicapa (MLPClassifier)**:

- **Capa de Entrada (TF-IDF Vector Space):**
  - Extracción de características textuales mediante `TfidfVectorizer` utilizando unigramas y bigramas (`ngram_range=(1, 2)`).
  - Ponderación sublineal de frecuencias (`sublinear_tf=True`) para suavizar el efecto de términos de alta frecuencia.
  - Dimensión del espacio vectorial de entrada: hasta 1000 términos y pares de palabras clave.
- **Capa Oculta (Extracción y Abstracción de Patrones):**
  - 32 neuronas artificiales densamente conectadas con función de activación no lineal **ReLU** ($f(x) = \\max(0, x)$).
  - Captura combinaciones no lineales de términos representativos (por ejemplo, cláusulas jurídicas, estructuras contables, términos de ingeniería de software o protocolos de gestión humana).
- **Capa de Salida:**
  - 6 neuronas de salida correspondientes a las categorías documentales del sistema, con función de activación **Softmax**, que genera una distribución probabilística normalizada:
    $$P(C_i | X) = \\frac{{e^{{z_i}}}}{{\\sum_{{j=1}}^{{K}} e^{{z_j}}}}$$
- **Algoritmo de Optimización y Entrenamiento:**
  - Optimizador **Adam** (Adaptive Moment Estimation) con tasa de aprendizaje inicial $\\eta = 0.005$.
  - Regularización de pesos L2 (penalty $\\alpha = 0.001$) para prevenir el sobreajuste.
  - Convergencia alcanzada en {evaluacion['n_iter']} épocas con una función de pérdida final de {evaluacion['loss_curve_final']:.4f}.

### Proceso de Entrenamiento y Validación
El dataset fue dividido de forma estratificada:
- **Datos de Entrenamiento:** {evaluacion['n_train']} documentos (80% del corpus etiquetado).
- **Datos de Prueba (Test Set):** {evaluacion['n_test']} documentos (20% del corpus independiente, nunca vistos por el modelo).

### Métricas de Validación Obtenidas
- **Exactitud Global (Accuracy):** **{acc_pct:.2f}%** en el conjunto de prueba independiente.

{tabla_metricas}

### Matriz de Confusión
{tabla_confusion}

---

## 3. Qué información registra la base de datos (Evidencia)

Para garantizar que las predicciones no queden únicamente en memoria RAM, se diseñó e implementó una estructura relacional en **SQLite** almacenada en el artefacto:  
`artifacts/evidencia_reconocimiento.db`.

### Esquema de la Tabla `evidencia_reconocimiento`
| Campo | Tipo | Propósito y Utilidad |
| :--- | :--- | :--- |
| `id` | `INTEGER PRIMARY KEY` | Identificador único autoincremental de la transacción o evidencia. |
| `timestamp` | `TEXT` | Marca de tiempo ISO-8601 en la que se ejecutó el reconocimiento. |
| `fuente_documento` | `TEXT` | Nombre del archivo analizado o identificador del documento. |
| `tipo_fuente` | `TEXT` | Formato físico o digital de origen (`pdf`, `docx`, `txt`, `texto_directo`). |
| `resumen_contenido` | `TEXT` | Extracto o muestra de los primeros caracteres del texto analizado. |
| `categoria_real` | `TEXT` | Etiqueta real conocida (o indicación de documento no etiquetado en producción). |
| `categoria_predicha` | `TEXT` | Tipo documental predicho por la red neuronal. |
| `nivel_confianza` | `REAL` | Probabilidad máxima arrojada por la capa Softmax de la RNA (0.00 a 1.00). |
| `estado_proceso` | `TEXT` | Indicador de control: `Procesado Exitosamente`, `Requiere Revisión` o `Rechazado`. |
| `categoria_dominio` | `TEXT` | Macro-área de la organización derivada de la ontología (`Dominio_Legal`, etc.). |
| `departamento_destino` | `TEXT` | Unidad administrativa responsable asignada por la ontología. |
| `accion_sugerida` | `TEXT` | Política operativa o procedimiento a aplicar sobre el documento. |
| `significado_ontologico` | `TEXT` | Frase interpretativa completa generada mediante el grafo del conocimiento. |

La base de datos cuenta con índices sobre `categoria_predicha` y `estado_proceso` para soportar consultas de auditoría, monitoreo de desviaciones e informes de gestión.

---

## 4. Qué conceptos y relaciones representa la ontología

La ontología del proyecto modela el conocimiento del dominio documental mediante un grafo dirigido estructurado, serializado en el estándar abierto:  
`artifacts/ontologia_documental.graphml`.

### Conceptos del Dominio (Nodos del Grafo)
1. **`DocumentoDigital`**: Representa el artefacto de entrada (archivo o texto).
2. **`RedNeuronal`**: Representa el agente de inteligencia artificial que procesa los patrones numéricos.
3. **`Prediccion`**: Representa la inferencia generada con su probabilidad asociada.
4. **`EvidenciaBD`**: Representa el registro persistente de auditoría en SQLite.
5. **`TipoDocumento`**: Representa la tipología clasificada (`contrato`, `factura`, `reporte_financiero`, etc.).
6. **`CategoriaDominio`**: Representa las áreas estratégicas de la institución (`Dominio_Legal`, `Dominio_Financiero`, `Dominio_Tecnologico`, `Dominio_Administrativo`).
7. **`DepartamentoResponsable`**: Representa las dependencias operativas (`Area_Juridica`, `Contabilidad_Tesoreria`, `Operaciones_TI`, `Gestion_Humana`).
8. **`AccionFlujo`**: Representa los protocolos de acción (`Custodia_Firmas_Legales`, `Auditoria_Tributaria`, `Despliegue_Documentacion_Tecnica`, `Difusion_Interna`).

### Relaciones entre Conceptos (Proposiciones Semánticas con Sentido de Frase)
Cada arista del grafo constituye una proposición semántica válida (`Sujeto -> Predicado -> Objeto`):

{relaciones_md}

### Interpretación del Significado
Cuando la red neuronal predice, por ejemplo, que un archivo corresponde a una `factura`, la ontología permite interpretar que dicho documento pertenece al `Dominio_Financiero`, debe enrutarse hacia `Contabilidad_Tesoreria` y exige la acción operativa de `Auditoria_Tributaria`. De este modo, la predicción matemática se transforma en **conocimiento accionable** para la organización.

---

## 5. Integración completa: Flujo Verificado

Se ejecutaron pruebas integradas con documentos de diversos tipos y procedencias en disco (`.txt`, `.docx`, `.md`) para comprobar el ciclo completo:
$$\\text{{Entrada}} \\longrightarrow \\text{{Red Neuronal (RNA)}} \\longrightarrow \\text{{Evidencia (SQLite)}} \\longrightarrow \\text{{Significado (Ontología)}}$$

### Evidencia de Registros Procesados
{tabla_casos}

---

## 6. Limitaciones encontradas

Durante la concepción, entrenamiento y validación del sistema se identificaron las siguientes limitaciones técnicas y operativas:

1. **Documentos basados en imágenes sin capa OCR:**  
   Si un archivo PDF o imagen contiene texto escaneado sin una capa OCR previa, la extracción textual produce cadenas vacías. El sistema cuenta con un mecanismo de detección temprana que rechaza el documento y sugiere un reescaneo con OCR, pero el clasificador neuronal en sí mismo no procesa píxeles directamente.
2. **Ambigüedad de vocabulario en documentos híbridos:**  
   Documentos como actas de junta directiva donde se discuten presupuestos financieros pueden contener términos simultáneos de finanzas y de gobernanza corporativa. En tales circunstancias, la red neuronal puede disminuir su nivel de certidumbre. El sistema mitiga este problema mediante un **umbral de confianza (threshold = 0.40)**: si la certeza es inferior, se clasifica en estado `Requiere Revisión Manual`.
3. **Escala y diversidad del corpus:**  
   El modelo actual fue entrenado sobre un conjunto de documentos estructurados en español. Textos provenientes de otras jurisdicciones o con jerga regional específica requerirán un proceso de reentrenamiento continuo (fine-tuning) alimentado por la misma base de datos de evidencia.

---

## 7. Estructura de Entregables en el Repositorio

El proyecto cumple a cabalidad con la estructura de archivos requerida:
```text
ia_semestre/
├── src/
│   └── semana08_representaciones.py       <- Código del modelo, base de datos y ontología
├── artifacts/
│   ├── modelo_red_neuronal.joblib        <- Red Neuronal MLP entrenada
│   ├── vectorizador_tfidf.joblib         <- Vocabulario y ponderaciones TF-IDF
│   ├── evidencia_reconocimiento.db       <- Base de datos SQLite con evidencia verificable
│   ├── ontologia_documental.graphml      <- Grafo de ontología en formato GraphML
│   └── muestras_prueba/                  <- Archivos reales (.txt, .docx, .md) procesados
└── reports/
    └── semana08.md                       <- Este informe técnico detallado
```
"""

    REPORT_PATH.write_text(contenido, encoding="utf-8")
    print(f"[+] Reporte técnico generado exitosamente en: {REPORT_PATH}")


# ==============================================================================
# 8. EJECUCIÓN PRINCIPAL
# ==============================================================================

def main():
    print("=" * 80)
    print("SISTEMA DE ANÁLISIS DOCUMENTAL - SEMANA 08: REPRESENTACIONES DEL RECONOCIMIENTO")
    print("=" * 80)

    # Paso 1: Inicializar Base de Datos SQLite
    print("\n[1/5] Inicializando Base de Datos SQLite de Evidencias...")
    inicializar_base_datos(DB_PATH)
    print(f"      Base de datos lista en: {DB_PATH}")

    # Paso 2: Construir y Exportar Ontología GraphML
    print("\n[2/5] Construyendo Ontología del Dominio y exportando a GraphML...")
    ontologia = construir_ontologia()
    guardar_ontologia_graphml(ontologia, ONTOLOGY_PATH)
    print(f"      Ontología guardada con {ontologia.number_of_nodes()} nodos y {ontologia.number_of_edges()} aristas en: {ONTOLOGY_PATH}")

    # Paso 3: Entrenar y Validar Red Neuronal Artificial
    print("\n[3/5] Entrenando Red Neuronal Artificial (MLPClassifier)...")
    modelo, vectorizador, evaluacion = entrenar_modelo_neuronal(random_state=42)
    print(f"      Entrenamiento completado en {evaluacion['n_iter']} épocas.")
    print(f"      Muestras de entrenamiento: {evaluacion['n_train']} | Muestras de prueba: {evaluacion['n_test']}")
    print(f"      Exactitud (Accuracy) obtenida: {evaluacion['accuracy'] * 100:.2f}%")

    # Paso 4: Demostración del Flujo Completo sobre Archivos de Prueba en Disco
    print("\n[4/5] Demostrando Flujo Completo: Entrada -> RNA -> Evidencia en BD -> Ontología...")
    muestras = crear_muestras_prueba()
    muestras_procesadas = []

    for ruta_archivo, cat_real in muestras:
        resultado = clasificar_y_registrar_documento(
            fuente=ruta_archivo,
            modelo=modelo,
            vectorizador=vectorizador,
            ontologia=ontologia,
            db_path=DB_PATH,
            categoria_real=cat_real,
            umbral_minimo=0.40
        )
        muestras_procesadas.append(resultado)

        print(f"\n-> Archivo: '{resultado['fuente_documento']}' [Formato: {resultado['tipo_fuente']}]")
        print(f"   * Resumen            : {resultado['resumen_contenido'][:75]}...")
        print(f"   * Categoría Real     : {cat_real}")
        print(f"   * Predicción RNA     : {resultado['categoria_predicha']} (Confianza: {resultado['nivel_confianza']*100:.2f}%)")
        print(f"   * Estado en BD       : Guardado con ID #{resultado['id_bd']} [{resultado['estado_proceso']}]")
        print(f"   * Dominio Ontológico : {resultado['categoria_dominio']}")
        print(f"   * Enrutamiento       : {resultado['departamento_destino']} -> Acción: {resultado['accion_sugerida']}")
        print(f"   * Significado        : {resultado['significado_ontologico']}")

    # Consultar evidencias en base de datos para verificar persistencia
    print("\n" + "-" * 80)
    print("VERIFICACIÓN DE EVIDENCIAS EN BASE DE DATOS SQLITE (Últimos registros)")
    print("-" * 80)
    evidencias_recientes = consultar_evidencias(DB_PATH, limite=7)
    for ev in evidencias_recientes:
        print(f"ID #{ev['id']:02d} | Doc: {ev['fuente_documento'][:30]:<30} | Fmt: {ev['tipo_fuente']:<5} | Pred: {ev['categoria_predicha']:<18} | Conf: {ev['nivel_confianza']*100:>5.1f}% | Depto: {ev['departamento_destino']}")

    # Consultar métricas agregadas de la BD
    stats_bd = consultar_estadisticas_bd(DB_PATH)
    print(f"\n[+] Total de registros históricos en BD: {stats_bd['total_registros']}")
    print(f"[+] Confianza promedio histórica: {stats_bd['confianza_promedio']*100:.2f}%")

    # Paso 5: Generar Informe Técnico Markdown
    print("\n[5/5] Generando reporte técnico detallado en reports/semana08.md...")
    generar_reporte_md(evaluacion, muestras_procesadas, ontologia)

    print("\n" + "=" * 80)
    print("FLUJO COMPLETADO EXITOSAMENTE:")
    print("  1. Red Neuronal MLP entrenada y guardada en artifacts/modelo_red_neuronal.joblib")
    print("  2. Vectorizador TF-IDF guardado en artifacts/vectorizador_tfidf.joblib")
    print("  3. Evidencia registrada y verificada en artifacts/evidencia_reconocimiento.db")
    print("  4. Ontología construida y serializada en artifacts/ontologia_documental.graphml")
    print("  5. Muestras reales (.txt, .docx, .md) creadas y evaluadas en artifacts/muestras_prueba/")
    print("  6. Reporte exhaustivo generado en reports/semana08.md")
    print("=" * 80)


if __name__ == "__main__":
    main()
